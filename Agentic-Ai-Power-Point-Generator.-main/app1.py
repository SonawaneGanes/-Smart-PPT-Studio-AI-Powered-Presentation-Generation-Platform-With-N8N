 
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.dml.color import RGBColor


# ---------------- THEMES ----------------

THEMES = {

    "blue": {
        "bg": RGBColor(15, 23, 42),
        "accent": RGBColor(59, 130, 246),
        "card": RGBColor(30, 41, 59),
        "text": RGBColor(241, 245, 249)
    },

    "green": {
        "bg": RGBColor(6, 20, 15),
        "accent": RGBColor(16, 185, 129),
        "card": RGBColor(20, 30, 25),
        "text": RGBColor(240, 253, 244)
    },

    "purple": {
        "bg": RGBColor(17, 24, 39),
        "accent": RGBColor(168, 85, 247),
        "card": RGBColor(31, 41, 55),
        "text": RGBColor(243, 244, 246)
    },

    "red": {
        "bg": RGBColor(35, 10, 10),
        "accent": RGBColor(239, 68, 68),
        "card": RGBColor(45, 20, 20),
        "text": RGBColor(255, 245, 245)
    }
}


# ---------------- PPT CLASS ----------------

class SmartPresentation:

    def __init__(self, topic, theme="blue"):

        self.topic = topic

        self.theme = THEMES.get(
            theme,
            THEMES["blue"]
        )

        self.prs = Presentation()

    # ---------------- BACKGROUND ----------------

    def set_background(self, slide):

        background = slide.background

        fill = background.fill

        fill.solid()

        fill.fore_color.rgb = self.theme["bg"]

    # ---------------- HEADER ----------------

    def add_header(self, slide, text):

        title_box = slide.shapes.add_textbox(
            Inches(0.5),
            Inches(0.4),
            Inches(8.5),
            Inches(0.8)
        )

        tf = title_box.text_frame

        p = tf.paragraphs[0]

        p.text = text

        p.font.size = Pt(28)

        p.font.bold = True

        p.font.color.rgb = self.theme["text"]

    # ---------------- CONTENT CARD ----------------

    def add_content_card(self, slide, points):

        card = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            Inches(0.7),
            Inches(1.5),
            Inches(8.2),
            Inches(4.6)
        )

        fill = card.fill

        fill.solid()

        fill.fore_color.rgb = self.theme["card"]

        line = card.line

        line.color.rgb = self.theme["accent"]

        line.width = Pt(2)

        tf = card.text_frame

        tf.clear()

        first = True

        for point in points:

            if first:

                p = tf.paragraphs[0]

                first = False

            else:

                p = tf.add_paragraph()

            p.text = f"• {point}"

            p.font.size = Pt(20)

            p.font.color.rgb = self.theme["text"]

            p.space_after = Pt(14)

    # ---------------- FOOTER ----------------

    def add_footer(self, slide):

        footer = slide.shapes.add_textbox(
            Inches(0.4),
            Inches(6.9),
            Inches(4),
            Inches(0.3)
        )

        tf = footer.text_frame

        p = tf.paragraphs[0]

        p.text = "Generated using Smart PPT Studio"

        p.font.size = Pt(10)

        p.font.color.rgb = RGBColor(180, 180, 180)

    # ---------------- TITLE SLIDE ----------------

    def create_title_slide(self):

        slide_layout = self.prs.slide_layouts[6]

        slide = self.prs.slides.add_slide(
            slide_layout
        )

        self.set_background(slide)

        title_box = slide.shapes.add_textbox(
            Inches(0.8),
            Inches(2),
            Inches(8),
            Inches(1)
        )

        tf = title_box.text_frame

        p = tf.paragraphs[0]

        p.text = self.topic

        p.font.size = Pt(36)

        p.font.bold = True

        p.font.color.rgb = self.theme["accent"]

        subtitle = slide.shapes.add_textbox(
            Inches(1),
            Inches(3),
            Inches(7),
            Inches(0.5)
        )

        tf2 = subtitle.text_frame

        p2 = tf2.paragraphs[0]

        p2.text = "AI Powered Professional Presentation"

        p2.font.size = Pt(18)

        p2.font.italic = True

        p2.font.color.rgb = self.theme["text"]

        self.add_footer(slide)

    # ---------------- NORMAL SLIDE ----------------

    def create_slide(self, title, points):

        slide_layout = self.prs.slide_layouts[6]

        slide = self.prs.slides.add_slide(
            slide_layout
        )

        self.set_background(slide)

        self.add_header(slide, title)

        self.add_content_card(slide, points)

        self.add_footer(slide)

    # ---------------- GENERATE PPT ----------------

    def generate(self):

        self.create_title_slide()

        slides = [

            (
                f"Introduction to {self.topic}",
                [
                    "Overview of the topic",
                    "Importance in modern industries",
                    "Core concepts and fundamentals",
                    "Emerging innovations and trends"
                ]
            ),

            (
                "Major Benefits",
                [
                    "Enhanced operational efficiency",
                    "Better decision making",
                    "Scalable implementation",
                    "Improved user experience"
                ]
            ),

            (
                "Challenges",
                [
                    "Complex implementation process",
                    "High infrastructure costs",
                    "Security and compliance risks",
                    "Requirement of skilled professionals"
                ]
            ),

            (
                "Industry Applications",
                [
                    "Healthcare and diagnostics",
                    "Finance and analytics",
                    "Retail and personalization",
                    "Manufacturing automation"
                ]
            ),

            (
                "Future Opportunities",
                [
                    "Rapid technological growth",
                    "AI driven innovation",
                    "Expansion into new industries",
                    "Increased automation adoption"
                ]
            ),

            (
                "Conclusion",
                [
                    "Technology adoption is accelerating",
                    "Businesses must adapt strategically",
                    "Innovation drives long term success",
                    "Future potential remains enormous"
                ]
            )

        ]

        for title, points in slides:

            self.create_slide(
                title,
                points
            )

        self.prs.save(
            "modern_presentation.pptx"
        )

        print(
            "Presentation created successfully."
        )


# ---------------- MAIN ----------------

if __name__ == "__main__":

    topic = {{TOPIC}}

    theme = "{{THEME}}"

    ppt = SmartPresentation(
        topic=topic,
        theme=theme
    )

    ppt.generate()

