from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.shared import Inches

# Create a presentation object
prs = Presentation()

# --- Title Slide ---
title_slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(title_slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]

title.text = "The Future of Sustainable Energy"
subtitle.text = "Innovations and Opportunities"

# --- Content Slide 1: Introduction to Sustainable Energy ---
content_slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(content_slide_layout)
title = slide.shapes.title
body_shape = slide.shapes.placeholders[1]
tf = body_shape.text_frame

title.text = "What is Sustainable Energy?"

p = tf.add_paragraph()
p.text = "Energy derived from natural resources that are replenished at a rate faster than they are consumed."
p.level = 0
p.font.size = Pt(20)
p.font.color.rgb = RGBColor(50, 50, 50)

p = tf.add_paragraph()
p.text = "Key characteristics:"
p.level = 0
p.font.size = Pt(20)
p.font.color.rgb = RGBColor(50, 50, 50)

p = tf.add_paragraph()
p.text = "Renewable and inexhaustible"
p.level = 1
p.font.size = Pt(18)
p.font.color.rgb = RGBColor(80, 80, 80)

p = tf.add_paragraph()
p.text = "Environmentally friendly (low or no emissions)"
p.level = 1
p.font.size = Pt(18)
p.font.color.rgb = RGBColor(80, 80, 80)

p = tf.add_paragraph()
p.text = "Cost-effective in the long run"
p.level = 1
p.font.size = Pt(18)
p.font.color.rgb = RGBColor(80, 80, 80)


# --- Content Slide 2: Major Types of Sustainable Energy ---
slide = prs.slides.add_slide(content_slide_layout)
title = slide.shapes.title
body_shape = slide.shapes.placeholders[1]
tf = body_shape.text_frame

title.text = "Key Sustainable Energy Sources"

p = tf.add_paragraph()
p.text = "Solar Energy: Harnessing sunlight through photovoltaic panels or concentrated solar power."
p.level = 0
p.font.size = Pt(18)
p.font.color.rgb = RGBColor(50, 50, 50)

p = tf.add_paragraph()
p.text = "Wind Energy: Utilizing wind turbines to convert wind's kinetic energy into electricity."
p.level = 0
p.font.size = Pt(18)
p.font.color.rgb = RGBColor(50, 50, 50)

p = tf.add_paragraph()
p.text = "Hydropower: Generating electricity from the movement of water, typically through dams."
p.level = 0
p.font.size = Pt(18)
p.font.color.rgb = RGBColor(50, 50, 50)

p = tf.add_paragraph()
p.text = "Geothermal Energy: Tapping into the Earth's internal heat for heating and electricity generation."
p.level = 0
p.font.size = Pt(18)
p.font.color.rgb = RGBColor(50, 50, 50)

p = tf.add_paragraph()
p.text = "Biomass Energy: Generating energy from organic matter like plants and animal waste."
p.level = 0
p.font.size = Pt(18)
p.font.color.rgb = RGBColor(50, 50, 50)


# --- Content Slide 3: Technological Advancements ---
slide = prs.slides.add_slide(content_slide_layout)
title = slide.shapes.title
body_shape = slide.shapes.placeholders[1]
tf = body_shape.text_frame

title.text = "Innovations Driving the Future"

p = tf.add_paragraph()
p.text = "Advanced Solar Technologies: Perovskite solar cells, bifacial panels, and floating solar farms."
p.level = 0
p.font.size = Pt(18)
p.font.color.rgb = RGBColor(50, 50, 50)

p = tf.add_paragraph()
p.text = "Next-Generation Wind Turbines: Larger, more efficient designs and offshore wind farms."
p.level = 0
p.font.size = Pt(18)
p.font.color.rgb = RGBColor(50, 50, 50)

p = tf.add_paragraph()
p.text = "Energy Storage Solutions: Lithium-ion batteries, flow batteries, and hydrogen storage."
p.level = 0
p.font.size = Pt(18)
p.font.color.rgb = RGBColor(50, 50, 50)

p = tf.add_paragraph()
p.text = "Smart Grids: Integrating renewable energy sources with advanced grid management systems."
p.level = 0
p.font.size = Pt(18)
p.font.color.rgb = RGBColor(50, 50, 50)

p = tf.add_paragraph()
p.text = "Artificial Intelligence (AI) in Energy: Predictive maintenance, demand forecasting, and grid optimization."
p.level = 0
p.font.size = Pt(18)
p.font.color.rgb = RGBColor(50, 50, 50)


# --- Content Slide 4: Benefits and Challenges ---
slide = prs.slides.add_slide(content_slide_layout)
title = slide.shapes.title
body_shape = slide.shapes.placeholders[1]
tf = body_shape.text_frame

title.text = "The Upside and the Hurdles"

p = tf.add_paragraph()
p.text = "Benefits:"
p.level = 0
p.font.size = Pt(20)
p.font.color.rgb = RGBColor(50, 50, 50)

p = tf.add_paragraph()
p.text = "Environmental Protection: Reduced greenhouse gas emissions and air pollution."
p.level = 1
p.font.size = Pt(18)
p.font.color.rgb = RGBColor(80, 80, 80)

p = tf.add_paragraph()
p.text = "Energy Independence: Reduced reliance on fossil fuels and volatile global markets."
p.level = 1
p.font.size = Pt(18)
p.font.color.rgb = RGBColor(80, 80, 80)

p = tf.add_paragraph()
p.text = "Economic Growth: Job creation in manufacturing, installation, and maintenance."
p.level = 1
p.font.size = Pt(18)
p.font.color.rgb = RGBColor(80, 80, 80)

p = tf.add_paragraph()
p.text = "Challenges:"
p.level = 0
p.font.size = Pt(20)
p.font.color.rgb = RGBColor(50, 50, 50)

p = tf.add_paragraph()
p.text = "Intermittency: Variability in solar and wind power generation."
p.level = 1
p.font.size = Pt(18)
p.font.color.rgb = RGBColor(80, 80, 80)

p = tf.add_paragraph()
p.text = "High Upfront Costs: Initial investment for renewable energy infrastructure."
p.level = 1
p.font.size = Pt(18)
p.font.color.rgb = RGBColor(80, 80, 80)

p = tf.add_paragraph()
p.text = "Grid Integration: Adapting existing grids to accommodate distributed renewable sources."
p.level = 1
p.font.size = Pt(18)
p.font.color.rgb = RGBColor(80, 80, 80)

p = tf.add_paragraph()
p.text = "Land Use and Environmental Impact: Considerations for large-scale installations."
p.level = 1
p.font.size = Pt(18)
p.font.color.rgb = RGBColor(80, 80, 80)


# --- Content Slide 5: The Path Forward ---
slide = prs.slides.add_slide(content_slide_layout)
title = slide.shapes.title
body_shape = slide.shapes.placeholders[1]
tf = body_shape.text_frame

title.text = "Moving Towards a Sustainable Future"

p = tf.add_paragraph()
p.text = "Policy and Investment: Government incentives, carbon pricing, and private sector funding."
p.level = 0
p.font.size = Pt(18)
p.font.color.rgb = RGBColor(50, 50, 50)

p = tf.add_paragraph()
p.text = "Research and Development: Continued innovation in efficiency and cost reduction."
p.level = 0
p.font.size = Pt(18)
p.font.color.rgb = RGBColor(50, 50, 50)

p = tf.add_paragraph()
p.text = "Public Awareness and Education: Fostering understanding and support for sustainable energy."
p.level = 0
p.font.size = Pt(18)
p.font.color.rgb = RGBColor(50, 50, 50)

p = tf.add_paragraph()
p.text = "International Cooperation: Sharing best practices and technologies globally."
p.level = 0
p.font.size = Pt(18)
p.font.color.rgb = RGBColor(50, 50, 50)

p = tf.add_paragraph()
p.text = "The transition to sustainable energy is not just an environmental imperative, but an economic and social opportunity."
p.level = 0
p.font.size = Pt(18)
p.font.color.rgb = RGBColor(50, 50, 50)


# --- Save the presentation ---
prs.save("presentation.pptx")